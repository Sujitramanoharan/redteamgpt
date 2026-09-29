"""Scan uploaded documents for indirect prompt injection.

An attacker does not need the user to type anything malicious. They hide the
instruction inside a file - a CV, an invoice, a report - and wait for an AI
system to ingest it. The user sees an ordinary document; the model sees an
instruction. This finds the hidden passage and points at it.
"""
import io
import logging
import re
from pathlib import Path
from typing import Optional

from config import settings
from firewall import detect

logger = logging.getLogger(__name__)

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".csv"}

# Bounds so a malicious upload cannot exhaust the server.
MAX_PAGES = 60
MAX_CHARS = 400_000
MAX_PASSAGES = 120
PASSAGE_TARGET = 700


class DocumentError(ValueError):
    """Raised when a document cannot be read or is not supported."""


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(f"Could not read PDF: {exc}") from exc

    if reader.is_encrypted:
        raise DocumentError("This PDF is password protected.")

    parts = []
    for page in reader.pages[:MAX_PAGES]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            # A single malformed page should not fail the whole scan.
            continue
    return "\n\n".join(parts)


def _extract_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(f"Could not read Word document: {exc}") from exc

    parts = [p.text for p in document.paragraphs]
    # Instructions are often parked inside tables, where readers skim past them.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise DocumentError(
            f"Unsupported file type '{suffix or 'unknown'}'. "
            f"Supported: {', '.join(sorted(SUPPORTED))}"
        )

    if suffix == ".pdf":
        text = _extract_pdf(data)
    elif suffix == ".docx":
        text = _extract_docx(data)
    else:
        text = data.decode("utf-8", errors="replace")

    if not text.strip():
        raise DocumentError(
            "No readable text found. Scanned images are not supported without OCR."
        )
    return text[:MAX_CHARS]


def split_passages(text: str) -> list[str]:
    """Split on blank lines, then pack into passages of roughly equal size.

    Scanning passage by passage is what lets the result name the offending
    paragraph instead of only judging the file as a whole.
    """
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    passages, current = [], ""

    for block in blocks:
        if len(current) + len(block) + 2 <= PASSAGE_TARGET:
            current = f"{current}\n\n{block}" if current else block
            continue
        if current:
            passages.append(current)
        # A single oversized block still has to be broken up.
        while len(block) > PASSAGE_TARGET:
            passages.append(block[:PASSAGE_TARGET])
            block = block[PASSAGE_TARGET:]
        current = block

    if current:
        passages.append(current)
    return passages[:MAX_PASSAGES]


def _offending_lines(passage: str, tokens: list) -> str:
    """Pull out the exact lines a signature matched.

    A passage can be several paragraphs, so quoting the whole thing back tells
    the reader nothing. The matched offsets identify the real line.
    """
    if not tokens:
        return ""

    lines, cursor, hits = passage.splitlines(keepends=True), 0, []
    spans = [(t["start"], t["end"]) for t in tokens if t.get("end", 0) > 0]

    for line in lines:
        start, end = cursor, cursor + len(line)
        if any(s < end and e > start for s, e in spans):
            hits.append(line.strip())
        cursor = end

    return "\n".join(h for h in hits if h)


def scan_document(filename: str, data: bytes, threshold: Optional[float] = None) -> dict:
    """Extract, split and scan a document, reporting the offending passages."""
    text = extract_text(filename, data)
    passages = split_passages(text)

    findings, highest = [], None
    for index, passage in enumerate(passages):
        result = detect(passage, threshold)
        if result["malicious"]:
            offending = _offending_lines(passage, result["detected_tokens"])
            findings.append({
                "passage_index": index,
                "excerpt": passage[:400],
                # The specific line to show the user, when we can pinpoint it.
                "offending_text": offending or passage[:200],
                "pinpointed": bool(offending),
                "verdict": result["verdict"],
                "risk_score": result["risk_score"],
                "category": result["category"],
                "priority": result["priority"],
                "explanation": result["explanation"],
                "detected_tokens": result["detected_tokens"],
            })
        if highest is None or result["risk_score"] > highest["risk_score"]:
            highest = result

    malicious = bool(findings)
    top = max(findings, key=lambda f: f["risk_score"]) if findings else None

    return {
        "filename": filename,
        "verdict": "BLOCKED" if malicious else "ALLOWED",
        "malicious": malicious,
        "risk_score": top["risk_score"] if top else (highest["risk_score"] if highest else 0),
        "priority": top["priority"] if top else (highest["priority"] if highest else None),
        "category": top["category"] if top else "Clean Document",
        "summary": (
            f"Found {len(findings)} suspicious passage"
            f"{'s' if len(findings) != 1 else ''} out of {len(passages)} scanned. "
            "An instruction hidden in a document can hijack any AI system that "
            "reads it, even though the person who uploaded it typed nothing."
            if malicious else
            f"Scanned {len(passages)} passages. No hidden instructions or "
            "injection attempts were found."
        ),
        "stats": {
            "characters": len(text),
            "passages_scanned": len(passages),
            "suspicious_passages": len(findings),
            "truncated": len(text) >= MAX_CHARS,
        },
        "findings": findings[:20],
        "preview": text[:1500],
    }
